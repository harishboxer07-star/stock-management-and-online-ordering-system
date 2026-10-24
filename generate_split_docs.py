import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import shutil

PROJECT_ROOT = r"c:\Users\DHASWEEN\OneDrive\Desktop\smart-grocery-ai-production project"
DESKTOP_DIR = r"c:\Users\DHASWEEN\OneDrive\Desktop"

# Color palette
NAVY_PRIMARY = RGBColor(30, 58, 138)      # #1E3A8A
BLUE_ACCENT = RGBColor(37, 99, 235)       # #2563EB
EMERALD_GREEN = RGBColor(5, 150, 105)     # #059669
DARK_TEXT = RGBColor(15, 23, 42)          # #0F172A
SLATE_GRAY = RGBColor(100, 116, 139)      # #64748B
WHITE = RGBColor(255, 255, 255)

LIGHT_BG_HEX = "F8FAFC"
META_BG_HEX = "F1F5F9"
ACCENT_BLUE_HEX = "2563EB"
ACCENT_GREEN_HEX = "059669"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", left_sz="4", other_sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{other_sz}" w:space="0" w:color="{top}"/>
            <w:left w:val="single" w:sz="{left_sz}" w:space="0" w:color="{left}"/>
            <w:bottom w:val="single" w:sz="{other_sz}" w:space="0" w:color="{bottom}"/>
            <w:right w:val="single" w:sz="{other_sz}" w:space="0" w:color="{right}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_field(p, field_name):
    run = p.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), field_name)
    run._r.append(fld)

def sanitize_file_content(path, content):
    if path.endswith('.env'):
        lines = []
        for line in content.splitlines():
            if 'MONGODB_URI=mongodb+srv://' in line:
                lines.append('# MONGODB_URI=mongodb+srv://<USERNAME>:<PASSWORD>@cluster0.mongodb.net/smart_grocery_db?retryWrites=true&w=majority')
            elif line.startswith('JWT_SECRET='):
                lines.append('JWT_SECRET=smart_grocery_super_secret_jwt_key_2026')
            else:
                lines.append(line)
        return '\n'.join(lines)
    return content

# BACKEND SPECIFIC SECTIONS
BACKEND_SECTIONS = [
    {
        "title": "Backend Section 1: Server Entry, Configuration & Seeds",
        "description": "Main Express HTTP application server, Socket.IO real-time event bus, database connection management, and demo database seeding.",
        "files": [
            {
                "path": "backend/package.json",
                "lang": "JSON",
                "role": "Backend Dependencies Manifest",
                "desc": "Declares backend libraries: Express, Mongoose, Socket.IO, PDFKit, JWT, bcryptjs, and Google Gemini AI SDK."
            },
            {
                "path": "backend/.env.example",
                "lang": "Env Configuration",
                "role": "Environment Variables Blueprint",
                "desc": "Configuration template for PORT, MONGODB_URI, JWT_SECRET, and GEMINI_API_KEY."
            },
            {
                "path": "backend/Dockerfile",
                "lang": "Dockerfile",
                "role": "Backend Containerfile",
                "desc": "Alpine Node.js container build configuration for isolated backend microservice hosting."
            },
            {
                "path": "backend/.dockerignore",
                "lang": "Docker Config",
                "role": "Docker Build Exclusions",
                "desc": "Prevents local development caches and modules from inflating container build contexts."
            },
            {
                "path": "backend/src/server.js",
                "lang": "JavaScript (Node.js)",
                "role": "Main Server Entry & WebSockets",
                "desc": "Primary HTTP application server initializing Express, CORS, Socket.IO server, WebSocket real-time broadcast hub, and REST route mounting."
            },
            {
                "path": "backend/src/config/db.js",
                "lang": "JavaScript (Node.js)",
                "role": "Database Connection & Seeder",
                "desc": "Mongoose connection manager handling MongoDB Atlas connection states, reconnect logic, and automated catalog seeding."
            },
            {
                "path": "backend/src/seed/seedData.js",
                "lang": "JavaScript (Node.js)",
                "role": "Database Seeding Script",
                "desc": "Pre-populates MongoDB with demo personas (Admin, Cashier, Delivery, Customer), catalog products, FEFO batches with expiry dates, and sample orders."
            }
        ]
    },
    {
        "title": "Backend Section 2: Data Models (Mongoose Schemas)",
        "description": "Object Data Modeling (ODM) schemas enforcing data validation, relationships, indexes, and virtual fields across all business domains.",
        "files": [
            {
                "path": "backend/src/models/User.js",
                "lang": "JavaScript (Mongoose)",
                "role": "User & Persona Schema",
                "desc": "User entity supporting roles ('admin', 'staff', 'delivery', 'customer'), bcrypt password hashing, and credentials verification."
            },
            {
                "path": "backend/src/models/Product.js",
                "lang": "JavaScript (Mongoose)",
                "role": "Master Product Catalog Schema",
                "desc": "Core product record storing SKU, name, category, base price, unit, reorder thresholds, and shelf life."
            },
            {
                "path": "backend/src/models/Batch.js",
                "lang": "JavaScript (Mongoose)",
                "role": "FEFO Inventory Batch Schema",
                "desc": "Tracks batch lot number, manufacturing date, expiry date, current stock, barcode (EAN/Code-128), and dynamic clearance markdown percentage."
            },
            {
                "path": "backend/src/models/Order.js",
                "lang": "JavaScript (Mongoose)",
                "role": "Sales Order & POS Receipt Schema",
                "desc": "Orders collection handling line items, batch-level allocations, taxes, discount deductions, payment method, and delivery status."
            },
            {
                "path": "backend/src/models/PurchaseOrder.js",
                "lang": "JavaScript (Mongoose)",
                "role": "Supplier Purchase Order Schema",
                "desc": "Procurement schema linking supplier IDs, requested items, expected delivery dates, and receiving confirmation statuses."
            },
            {
                "path": "backend/src/models/Supplier.js",
                "lang": "JavaScript (Mongoose)",
                "role": "Vendor Registry Schema",
                "desc": "Supplier registry storing company details, contact person, email, phone, product categories supplied, and lead time."
            },
            {
                "path": "backend/src/models/DailySale.js",
                "lang": "JavaScript (Mongoose)",
                "role": "Historical Sales Aggregations",
                "desc": "Aggregated daily product sales records utilized by the AI forecasting engine for demand trend modeling."
            },
            {
                "path": "backend/src/models/AuditLog.js",
                "lang": "JavaScript (Mongoose)",
                "role": "Immutable Security & Audit Trail",
                "desc": "Tamper-evident audit log storing timestamps, user ID, role, action category, affected resources, client IP, and diffs."
            }
        ]
    },
    {
        "title": "Backend Section 3: Security & Middleware",
        "description": "Request authentication guards, role-based authorization boundaries, and automated audit logging interceptors.",
        "files": [
            {
                "path": "backend/src/middleware/auth.js",
                "lang": "JavaScript (Node.js)",
                "role": "JWT Authentication & RBAC Guard",
                "desc": "Validates Authorization Bearer tokens, decrypts JWT payloads, attaches user records, and enforces role permissions."
            },
            {
                "path": "backend/src/middleware/audit.js",
                "lang": "JavaScript (Node.js)",
                "role": "Audit Logging Middleware",
                "desc": "Express interceptor that captures high-impact state mutations (POST, PUT, DELETE) and records compliance audit entries."
            }
        ]
    },
    {
        "title": "Backend Section 4: Business Logic & Core Services",
        "description": "Algorithmic engines: Google Gemini AI integration, heuristic demand forecasting, FEFO batch allocation with clearance markdowns, and PDF invoice generation.",
        "files": [
            {
                "path": "backend/src/services/aiService.js",
                "lang": "JavaScript (Node.js)",
                "role": "Heuristic Simulation AI Engine",
                "desc": "Offline fallback engine providing statistical moving-average demand forecasting, recipe generation, and executive summaries."
            },
            {
                "path": "backend/src/services/geminiService.js",
                "lang": "JavaScript (Node.js)",
                "role": "Google Gemini Generative AI SDK",
                "desc": "Integrates with Google Gemini API for conversational assistance, cart recipe suggestions, and supply-chain executive briefs."
            },
            {
                "path": "backend/src/services/inventoryService.js",
                "lang": "JavaScript (Node.js)",
                "role": "FEFO & Markdown Engine",
                "desc": "Calculates automated tiered markdowns (20%-50% discount) for batches nearing expiry and allocates stock First-Expired, First-Out."
            },
            {
                "path": "backend/src/services/invoiceService.js",
                "lang": "JavaScript (Node.js)",
                "role": "PDF Tax Invoice Generator",
                "desc": "Compiles formatted PDF tax invoices with store branding, cashier details, itemized batch numbers, GST, and totals."
            }
        ]
    },
    {
        "title": "Backend Section 5: API Controllers",
        "description": "Request dispatchers orchestrating business services, database queries, and response payloads across all backend domains.",
        "files": [
            {
                "path": "backend/src/controllers/authController.js",
                "lang": "JavaScript (Node.js)",
                "role": "Authentication Controller",
                "desc": "Handles customer registration, credential verification, JWT token signing, and persona profile retrieval."
            },
            {
                "path": "backend/src/controllers/productController.js",
                "lang": "JavaScript (Node.js)",
                "role": "Product Catalog Controller",
                "desc": "Manages product catalog CRUD, category queries, search indexing, and real-time aggregate stock calculation."
            },
            {
                "path": "backend/src/controllers/orderController.js",
                "lang": "JavaScript (Node.js)",
                "role": "Order Processing & POS Controller",
                "desc": "Executes atomic order creation with FEFO batch stock deduction, handles payments, triggers PDF invoices, and updates deliveries."
            },
            {
                "path": "backend/src/controllers/batchController.js",
                "lang": "JavaScript (Node.js)",
                "role": "FEFO Batch Inventory Controller",
                "desc": "Handles stock receiving, barcode lookups, expiry inspection, clearance markdown adjustments, and batch disposal."
            },
            {
                "path": "backend/src/controllers/reorderController.js",
                "lang": "JavaScript (Node.js)",
                "role": "Smart Reorder Controller",
                "desc": "Computes recommended reorder quantities based on stock thresholds and creates official Purchase Orders for fulfillment."
            },
            {
                "path": "backend/src/controllers/supplierController.js",
                "lang": "JavaScript (Node.js)",
                "role": "Supplier Registry Controller",
                "desc": "Manages vendor records, supplier performance metrics, contact details, and catalog category affiliations."
            },
            {
                "path": "backend/src/controllers/aiController.js",
                "lang": "JavaScript (Node.js)",
                "role": "AI Features Controller",
                "desc": "Connects frontend with Google Gemini AI and heuristic models for chat assistance, recipe generation, and 7-day forecasting."
            },
            {
                "path": "backend/src/controllers/analyticsController.js",
                "lang": "JavaScript (Node.js)",
                "role": "Executive Analytics Controller",
                "desc": "Aggregates key business metrics including total revenue, gross margin, at-risk expiring inventory, top SKUs, and sales trends."
            },
            {
                "path": "backend/src/controllers/auditController.js",
                "lang": "JavaScript (Node.js)",
                "role": "Audit Log Controller",
                "desc": "Serves filtered audit logs for compliance review by role, date range, action type, and severity."
            }
        ]
    },
    {
        "title": "Backend Section 6: API Route Definitions",
        "description": "Express router manifests connecting HTTP endpoints with corresponding controller actions and security middleware.",
        "files": [
            {
                "path": "backend/src/routes/authRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Auth Route Manifest",
                "desc": "Defines endpoints for user registration (/register), login (/login), and user profile verification (/me)."
            },
            {
                "path": "backend/src/routes/productRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Product Route Manifest",
                "desc": "Defines public catalog endpoints and protected staff/admin product modification routes."
            },
            {
                "path": "backend/src/routes/orderRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Order Route Manifest",
                "desc": "Defines order submission, POS checkout, invoice PDF streaming, and delivery status updates."
            },
            {
                "path": "backend/src/routes/batchRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Batch Route Manifest",
                "desc": "Defines FEFO batch inwarding, barcode lookup (/barcode/:code), and markdown update routes."
            },
            {
                "path": "backend/src/routes/reorderRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Reorder & PO Route Manifest",
                "desc": "Defines routes for automated reorder recommendations and purchase order status transitions."
            },
            {
                "path": "backend/src/routes/supplierRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Supplier Route Manifest",
                "desc": "Defines vendor registry CRUD endpoints accessible by staff and administrators."
            },
            {
                "path": "backend/src/routes/aiRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "AI Route Manifest",
                "desc": "Defines AI conversational assistant (/chat), recipe generator (/recipes), demand forecast (/forecast), and executive brief (/brief)."
            },
            {
                "path": "backend/src/routes/analyticsRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Analytics Route Manifest",
                "desc": "Defines executive dashboard analytical endpoints including revenue trends and wastage metrics."
            },
            {
                "path": "backend/src/routes/auditRoutes.js",
                "lang": "JavaScript (Express Router)",
                "role": "Audit Route Manifest",
                "desc": "Defines protected administrator endpoints for querying system audit trails."
            }
        ]
    }
]

# FRONTEND SPECIFIC SECTIONS
FRONTEND_SECTIONS = [
    {
        "title": "Frontend Section 1: Build Configuration & HTML Shell",
        "description": "Vite bundler configuration, Tailwind CSS design system tokens, PostCSS processing, NGINX reverse proxy, and single-page application HTML shell.",
        "files": [
            {
                "path": "frontend/package.json",
                "lang": "JSON",
                "role": "Frontend Dependencies Manifest",
                "desc": "Declares client dependencies including React 19, Lucide React icons, Axios, Socket.IO Client, Tailwind CSS, and Vite."
            },
            {
                "path": "frontend/vite.config.js",
                "lang": "JavaScript (Vite)",
                "role": "Vite Bundler Configuration",
                "desc": "Configures React fast-refresh, build optimizations, asset chunking, and local API proxy to the backend."
            },
            {
                "path": "frontend/tailwind.config.js",
                "lang": "JavaScript (Tailwind)",
                "role": "Tailwind Design System Tokens",
                "desc": "Custom theme configuration defining brand color palettes (emerald, violet, amber, slate), border-radii, and animations."
            },
            {
                "path": "frontend/postcss.config.js",
                "lang": "JavaScript (PostCSS)",
                "role": "PostCSS Pipeline",
                "desc": "Configures Tailwind CSS and Autoprefixer processors for cross-browser styling compatibility."
            },
            {
                "path": "frontend/index.html",
                "lang": "HTML5",
                "role": "Single-Page Application Shell",
                "desc": "HTML5 entry document importing Google Fonts (Inter, Plus Jakarta Sans), responsive viewport tags, and root React DOM node."
            },
            {
                "path": "frontend/vercel.json",
                "lang": "JSON",
                "role": "Vercel Edge Rewrite Rules",
                "desc": "Configures SPA rewrite fallback rules for seamless client-side routing on Vercel Edge networks."
            },
            {
                "path": "frontend/nginx.conf",
                "lang": "NGINX Config",
                "role": "Production NGINX Web Server",
                "desc": "Production NGINX web server configuration providing gzip compression, client-side route fallback to index.html, and API reverse proxying."
            },
            {
                "path": "frontend/Dockerfile",
                "lang": "Dockerfile",
                "role": "Frontend Production Containerfile",
                "desc": "Multi-stage Dockerfile that builds Vite static assets and serves them via a lightweight Alpine NGINX container."
            },
            {
                "path": "frontend/.gitignore",
                "lang": "Git Configuration",
                "role": "Frontend Version Control Exclusions",
                "desc": "Excludes frontend node_modules, dist build outputs, and local development caches."
            },
            {
                "path": "frontend/.dockerignore",
                "lang": "Docker Configuration",
                "role": "Frontend Docker Exclusions",
                "desc": "Excludes build caches and local node_modules from frontend Docker context."
            },
            {
                "path": "frontend/.oxlintrc.json",
                "lang": "JSON",
                "role": "Linter Configuration",
                "desc": "Configures fast Oxlint linter rules for code cleanliness and syntax validation."
            },
            {
                "path": "frontend/README.md",
                "lang": "Markdown",
                "role": "Frontend Documentation",
                "desc": "Client-specific setup instructions, available npm scripts, and environment variable guidelines."
            }
        ]
    },
    {
        "title": "Frontend Section 2: Main Application Entry & Global Styles",
        "description": "React root mounting, master application component with navigation shell, responsive layout wrappers, and global design styles.",
        "files": [
            {
                "path": "frontend/src/main.jsx",
                "lang": "JavaScript (React 19 JSX)",
                "role": "React 19 DOM Root Mount",
                "desc": "Mounts the React application into the root DOM element and wraps it with StrictMode and Context Providers."
            },
            {
                "path": "frontend/src/App.jsx",
                "lang": "JavaScript (React 19 JSX)",
                "role": "Master Application Shell & Router",
                "desc": "Central layout component featuring the navbar, quick persona switcher banner, modal controllers, and dynamic view routing."
            },
            {
                "path": "frontend/src/index.css",
                "lang": "CSS",
                "role": "Global Design Styles & Utilities",
                "desc": "Tailwind layer directives, custom utility classes, glassmorphism blur effects, and base typography."
            },
            {
                "path": "frontend/src/App.css",
                "lang": "CSS",
                "role": "Component Keyframe Animations",
                "desc": "Keyframe animations for pulse badges, scanner beam effects, AI glowing gradients, and modal transitions."
            }
        ]
    },
    {
        "title": "Frontend Section 3: State Management (Contexts)",
        "description": "Global React Context providers handling authentication state, user persona switching, shopping cart state, and real-time inventory validation.",
        "files": [
            {
                "path": "frontend/src/context/AuthContext.jsx",
                "lang": "JavaScript (React Context)",
                "role": "Authentication & Persona Context",
                "desc": "Manages global user session, JWT storage in localStorage, rapid 1-click persona switching (Admin, Staff, Delivery, Customer), and role helpers."
            },
            {
                "path": "frontend/src/context/CartContext.jsx",
                "lang": "JavaScript (React Context)",
                "role": "Shopping Cart & Stock Context",
                "desc": "Manages shopping cart line items, quantity adjustments, coupon application, stock limit validation, and local persistence."
            }
        ]
    },
    {
        "title": "Frontend Section 4: Shared Components & Modals",
        "description": "Reusable UI components: responsive navigation header, role-based access control gates, alerts drawer, hardware barcode scanner, and Gemini conversational assistant.",
        "files": [
            {
                "path": "frontend/src/components/Navbar.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Omnichannel Navigation Bar",
                "desc": "Top navigation component displaying brand identity, role-specific portal tabs, real-time stock alert badge, active persona badge, and shopping cart counter."
            },
            {
                "path": "frontend/src/components/RoleGate.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Declarative RBAC Guard Component",
                "desc": "Security wrapper component that checks user permissions against required roles before rendering protected views."
            },
            {
                "path": "frontend/src/components/AlertsDrawer.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Live Inventory Alerts Drawer",
                "desc": "Slide-out notification drawer displaying real-time WebSocket alerts for low stock levels and items nearing expiry."
            },
            {
                "path": "frontend/src/components/BarcodeScannerModal.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Interactive Barcode Scanner Modal",
                "desc": "Simulated hardware barcode scanner with visual laser scanning animation, quick-click barcode chips, and manual barcode entry."
            },
            {
                "path": "frontend/src/components/GeminiAssistantModal.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Gemini AI Assistant & Chef Modal",
                "desc": "Generative AI assistant modal supporting grocery shopping Q&A, dietary meal planning, and automated recipe generation from cart items."
            }
        ]
    },
    {
        "title": "Frontend Section 5: Application Pages & Portals",
        "description": "Customer Storefront, Cashier POS, Cart Checkout, FEFO Inventory, AI Demand Forecasting, Purchase Orders, Suppliers, Admin BI Dashboard, Audit Trail, and Delivery Portal.",
        "files": [
            {
                "path": "frontend/src/pages/Storefront.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Customer E-Commerce Catalog",
                "desc": "Customer shopping page featuring category filtering, live search, FEFO clearance discount badges, and one-click add to cart."
            },
            {
                "path": "frontend/src/pages/PosCounter.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "High-Speed Cashier POS Terminal",
                "desc": "In-store cashier point-of-sale interface featuring barcode SKU scanning, line-item adjustments, multi-mode payment, and instant PDF receipt generation."
            },
            {
                "path": "frontend/src/pages/CartCheckout.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Cart, Coupon & Checkout Portal",
                "desc": "Complete checkout experience with line item breakdown, coupon code discounts, address entry, simulated payment (UPI/Card/COD), and order confirmation."
            },
            {
                "path": "frontend/src/pages/InventoryManagement.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "FEFO Batch Inventory Controller",
                "desc": "Inventory management hub featuring batch receiving, expiry countdowns with warning badges (Fresh, Expiring Soon, Critical), and markdown controls."
            },
            {
                "path": "frontend/src/pages/AIForecasting.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "7-Day AI Demand Forecasting",
                "desc": "Predictive analytics presenting 7-day SKU demand projections, stockout probability estimates, trend charts, and automated reorder recommendations."
            },
            {
                "path": "frontend/src/pages/PurchaseOrders.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Supplier Purchase Order Hub",
                "desc": "Procurement portal for creating, approving, and tracking vendor purchase orders, with automated receiving workflow."
            },
            {
                "path": "frontend/src/pages/SuppliersPage.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Supplier Directory & Relations",
                "desc": "Vendor registry page managing supplier contacts, supported product categories, average delivery lead times, and reliability ratings."
            },
            {
                "path": "frontend/src/pages/AdminDashboard.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Executive BI Analytics & Gemini Brief",
                "desc": "Executive dashboard featuring revenue KPI cards, sales charts, expiry loss analytics, inventory health gauges, and Gemini Executive Briefs."
            },
            {
                "path": "frontend/src/pages/StaffAuditLogs.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Compliance & Security Audit Trail",
                "desc": "Audit investigation dashboard with real-time search, severity filtering, role categorization, and CSV export."
            },
            {
                "path": "frontend/src/pages/DeliveryPortal.jsx",
                "lang": "JavaScript (React JSX)",
                "role": "Delivery Driver Dispatch Portal",
                "desc": "Mobile-optimized delivery portal for drivers to view dispatch routes, customer delivery addresses, contact details, and update milestones."
            }
        ]
    },
    {
        "title": "Frontend Section 6: Client Services & Mock Fallbacks",
        "description": "Network communication layer: Axios HTTP client, Socket.IO client event handling, and comprehensive standalone mock data generator.",
        "files": [
            {
                "path": "frontend/src/services/api.js",
                "lang": "JavaScript",
                "role": "Centralized API & WebSocket Client",
                "desc": "Axios HTTP wrapper configuring base URLs, auth token request interceptors, error handling, and Socket.IO real-time subscriptions."
            },
            {
                "path": "frontend/src/services/mockFallback.js",
                "lang": "JavaScript",
                "role": "Autonomous Offline Mock Engine",
                "desc": "Extensive in-memory mock engine providing realistic fallback responses for products, FEFO batches, orders, AI forecasting, and analytics."
            }
        ]
    }
]

def generate_document(doc_title, doc_badge, doc_desc, sections, output_filename, accent_color_hex, accent_rgb):
    print(f"\n=======================================================")
    print(f"Building: {doc_title}")
    print(f"Output File: {output_filename}")
    print(f"=======================================================")

    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)

    section = doc.sections[0]
    section.different_first_page_header_footer = True

    # COVER PAGE
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(36)
    p_top.paragraph_format.space_after = Pt(12)

    p_badge = doc.add_paragraph()
    p_badge.paragraph_format.space_before = Pt(0)
    p_badge.paragraph_format.space_after = Pt(6)
    run_badge = p_badge.add_run(doc_badge)
    run_badge.font.name = "Segoe UI"
    run_badge.font.size = Pt(11)
    run_badge.font.bold = True
    run_badge.font.color.rgb = accent_rgb

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run(doc_title)
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(30)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run(doc_desc)
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = SLATE_GRAY

    # Accent decorative bar table
    bar_tbl = doc.add_table(rows=1, cols=1)
    bar_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    bar_cell = bar_tbl.cell(0, 0)
    bar_cell.width = Inches(7.0)
    set_cell_background(bar_cell, accent_color_hex)
    bar_p = bar_cell.paragraphs[0]
    bar_p.paragraph_format.space_before = Pt(2)
    bar_p.paragraph_format.space_after = Pt(2)
    bar_run = bar_p.add_run(" ")
    bar_run.font.size = Pt(2)

    doc.add_paragraph().paragraph_format.space_before = Pt(20)

    # Gather stats
    total_files = 0
    total_lines = 0
    file_entries = []

    for sec in sections:
        for f in sec["files"]:
            rel_path = f["path"]
            full_path = os.path.join(PROJECT_ROOT, rel_path)
            lines_cnt = 0
            size_kb = 0
            if os.path.exists(full_path):
                try:
                    with open(full_path, 'r', encoding='utf-8', errors='ignore') as fp:
                        lines_cnt = len(fp.readlines())
                        size_kb = round(os.path.getsize(full_path) / 1024, 1)
                except Exception:
                    lines_cnt = 0
            total_files += 1
            total_lines += lines_cnt
            file_entries.append({
                "path": rel_path,
                "role": f["role"],
                "lang": f["lang"],
                "lines": lines_cnt,
                "size_kb": size_kb
            })

    meta_items = [
        ("Platform", "SmartGrocery AI • Enterprise Retail & Quick-Commerce"),
        ("Module Scope", doc_title),
        ("Total Files", f"{total_files} Files Included"),
        ("Total Code Volume", f"{total_lines:,} Lines of Production Code"),
        ("Compiled Date", "September 2026"),
        ("Environment", "Node.js v24, React 19, Vite, MongoDB, Socket.IO, Gemini AI"),
        ("Confidentiality", "Internal Production Source Code Documentation")
    ]

    meta_tbl = doc.add_table(rows=len(meta_items), cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False

    for idx, (label, val) in enumerate(meta_items):
        c0 = meta_tbl.cell(idx, 0)
        c1 = meta_tbl.cell(idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)

        row_bg = META_BG_HEX if idx % 2 == 0 else "FFFFFF"
        set_cell_background(c0, row_bg)
        set_cell_background(c1, row_bg)
        set_cell_margins(c0, top=70, bottom=70, left=120, right=120)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
        set_cell_borders(c0, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0", other_sz="2")
        set_cell_borders(c1, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0", other_sz="2")

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Segoe UI"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = DARK_TEXT

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = DARK_TEXT

    doc.add_page_break()

    # HEADER / FOOTER
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"SmartGrocery AI • {doc_title}")
    hrun.font.name = "Segoe UI"
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGBColor(148, 163, 184)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    frun1 = fp.add_run(f"SmartGrocery AI Platform\t\tPage ")
    frun1.font.name = "Segoe UI"
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = SLATE_GRAY
    add_field(fp, "PAGE")
    frun2 = fp.add_run(" of ")
    frun2.font.name = "Segoe UI"
    frun2.font.size = Pt(8.5)
    frun2.font.color.rgb = SLATE_GRAY
    add_field(fp, "NUMPAGES")

    # FILE INDEX TABLE
    h_idx = doc.add_paragraph()
    h_idx.paragraph_format.space_before = Pt(12)
    h_idx.paragraph_format.space_after = Pt(4)
    r_idx_title = h_idx.add_run(f"{doc_title} — File Directory & Index")
    r_idx_title.font.name = "Segoe UI"
    r_idx_title.font.size = Pt(18)
    r_idx_title.font.bold = True
    r_idx_title.font.color.rgb = NAVY_PRIMARY

    idx_table = doc.add_table(rows=len(file_entries) + 1, cols=5)
    idx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    idx_table.autofit = False

    cols_width = [Inches(0.6), Inches(2.6), Inches(2.3), Inches(0.8), Inches(0.7)]
    idx_headers = ["#", "File Path", "Role / Component", "Language", "Lines"]

    for c_idx, htext in enumerate(idx_headers):
        cell = idx_table.cell(0, c_idx)
        cell.width = cols_width[c_idx]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
        p = cell.paragraphs[0]
        r = p.add_run(htext)
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = WHITE
        if c_idx in [0, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for idx, item in enumerate(file_entries, start=1):
        bg_col = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        row_cells = [
            str(idx),
            item["path"],
            item["role"],
            item["lang"],
            f'{item["lines"]:,}'
        ]
        for c_idx, val in enumerate(row_cells):
            cell = idx_table.cell(idx, c_idx)
            cell.width = cols_width[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            set_cell_borders(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0", other_sz="2")
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Consolas" if c_idx == 1 else "Segoe UI"
            r.font.size = Pt(7.5) if c_idx == 1 else Pt(8)
            r.font.color.rgb = DARK_TEXT
            if c_idx == 1:
                r.font.bold = True
            if c_idx in [0, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # SOURCE CODE LISTINGS
    file_counter = 0

    for sec in sections:
        doc.add_page_break()

        # Section Header Banner Table
        sec_banner = doc.add_table(rows=1, cols=1)
        sec_banner.alignment = WD_TABLE_ALIGNMENT.CENTER
        sec_banner.autofit = False
        s_cell = sec_banner.cell(0, 0)
        s_cell.width = Inches(7.0)
        set_cell_background(s_cell, "1E3A8A")
        set_cell_margins(s_cell, top=140, bottom=140, left=180, right=180)
        set_cell_borders(s_cell, top="1E3A8A", bottom="1E3A8A", left=accent_color_hex, right="1E3A8A", left_sz="36", other_sz="4")

        sp = s_cell.paragraphs[0]
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(2)
        srun_title = sp.add_run(sec["title"])
        srun_title.font.name = "Segoe UI"
        srun_title.font.size = Pt(15)
        srun_title.font.bold = True
        srun_title.font.color.rgb = WHITE

        sp_desc = s_cell.add_paragraph()
        sp_desc.paragraph_format.space_before = Pt(2)
        sp_desc.paragraph_format.space_after = Pt(0)
        srun_desc = sp_desc.add_run(sec["description"])
        srun_desc.font.name = "Segoe UI"
        srun_desc.font.size = Pt(9.5)
        srun_desc.font.color.rgb = RGBColor(226, 232, 240)

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        for f_info in sec["files"]:
            file_counter += 1
            rel_path = f_info["path"]
            full_path = os.path.join(PROJECT_ROOT, rel_path)

            content = ""
            lines = []
            if os.path.exists(full_path):
                try:
                    with open(full_path, 'r', encoding='utf-8', errors='ignore') as fp:
                        content = fp.read()
                        content = sanitize_file_content(rel_path, content)
                        lines = content.splitlines()
                except Exception as ex:
                    content = f"// Error reading file: {str(ex)}"
                    lines = [content]
            else:
                content = "// File not found at path"
                lines = [content]

            num_lines = len(lines)
            size_kb = round(len(content.encode('utf-8')) / 1024, 1)

            # Metadata Table
            f_meta_tbl = doc.add_table(rows=2, cols=1)
            f_meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            f_meta_tbl.autofit = False

            r0_cell = f_meta_tbl.cell(0, 0)
            r0_cell.width = Inches(7.0)
            set_cell_background(r0_cell, "0F172A")
            set_cell_margins(r0_cell, top=80, bottom=80, left=140, right=140)
            set_cell_borders(r0_cell, top="0F172A", bottom="1E293B", left=accent_color_hex, right="0F172A", left_sz="24", other_sz="4")

            p_fhead = r0_cell.paragraphs[0]
            p_fhead.paragraph_format.space_before = Pt(0)
            p_fhead.paragraph_format.space_after = Pt(0)
            
            rf_icon = p_fhead.add_run("📄 ")
            rf_icon.font.size = Pt(10)
            
            rf_path = p_fhead.add_run(rel_path)
            rf_path.font.name = "Consolas"
            rf_path.font.size = Pt(10.5)
            rf_path.font.bold = True
            rf_path.font.color.rgb = WHITE

            rf_meta = p_fhead.add_run(f"   |   {f_info['lang']}   •   {num_lines:,} lines ({size_kb} KB)")
            rf_meta.font.name = "Segoe UI"
            rf_meta.font.size = Pt(8.5)
            rf_meta.font.color.rgb = RGBColor(148, 163, 184)

            r1_cell = f_meta_tbl.cell(1, 0)
            r1_cell.width = Inches(7.0)
            set_cell_background(r1_cell, "F1F5F9")
            set_cell_margins(r1_cell, top=60, bottom=60, left=140, right=140)
            set_cell_borders(r1_cell, top="E2E8F0", bottom="CBD5E1", left=accent_color_hex, right="CBD5E1", left_sz="24", other_sz="2")

            p_fdesc = r1_cell.paragraphs[0]
            p_fdesc.paragraph_format.space_before = Pt(0)
            p_fdesc.paragraph_format.space_after = Pt(0)
            
            r_role = p_fdesc.add_run(f"{f_info['role']}: ")
            r_role.font.name = "Segoe UI"
            r_role.font.size = Pt(8.5)
            r_role.font.bold = True
            r_role.font.color.rgb = DARK_TEXT

            r_fbody = p_fdesc.add_run(f_info['desc'])
            r_fbody.font.name = "Segoe UI"
            r_fbody.font.size = Pt(8.5)
            r_fbody.font.color.rgb = DARK_TEXT

            doc.add_paragraph().paragraph_format.space_before = Pt(2)

            # Code container table
            code_tbl = doc.add_table(rows=1, cols=1)
            code_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            code_tbl.autofit = False

            code_cell = code_tbl.cell(0, 0)
            code_cell.width = Inches(7.0)
            set_cell_background(code_cell, "F8FAFC")
            set_cell_margins(code_cell, top=80, bottom=80, left=100, right=100)
            set_cell_borders(code_cell, top="E2E8F0", bottom="E2E8F0", left="3B82F6", right="E2E8F0", left_sz="16", other_sz="2")

            num_digits = max(3, len(str(num_lines)))

            for l_idx, line in enumerate(lines, start=1):
                p_code = code_cell.paragraphs[0] if l_idx == 1 else code_cell.add_paragraph()
                p_code.paragraph_format.space_before = Pt(0)
                p_code.paragraph_format.space_after = Pt(0)
                p_code.paragraph_format.line_spacing = 1.0

                r_num = p_code.add_run(f"{l_idx:>{num_digits}} | ")
                r_num.font.name = "Consolas"
                r_num.font.size = Pt(8.0)
                r_num.font.color.rgb = RGBColor(148, 163, 184)

                r_text = p_code.add_run(line)
                r_text.font.name = "Consolas"
                r_text.font.size = Pt(8.0)
                r_text.font.color.rgb = DARK_TEXT

            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_before = Pt(6)
            p_after.paragraph_format.space_after = Pt(14)

    target_project_path = os.path.join(PROJECT_ROOT, output_filename)
    target_desktop_path = os.path.join(DESKTOP_DIR, output_filename)

    doc.save(target_project_path)
    shutil.copy2(target_project_path, target_desktop_path)
    print(f"Successfully created: {target_desktop_path} ({os.path.getsize(target_desktop_path)} bytes)")

# Run both generations
generate_document(
    doc_title="SmartGrocery AI • Backend Source Code",
    doc_badge="NODE.JS REST API • MONGOOSE • SOCKET.IO • GEMINI AI",
    doc_desc="Complete Backend Architectural Blueprint, Database Models, Middleware, Services & API Controllers",
    sections=BACKEND_SECTIONS,
    output_filename="SmartGrocery_AI_Backend_Source_Code.docx",
    accent_color_hex=ACCENT_BLUE_HEX,
    accent_rgb=BLUE_ACCENT
)

generate_document(
    doc_title="SmartGrocery AI • Frontend Source Code",
    doc_badge="REACT 19 • VITE • TAILWIND CSS • SINGLE-PAGE APP",
    doc_desc="Complete Frontend Design System, Global Contexts, Reusable Components, Portals & Pages",
    sections=FRONTEND_SECTIONS,
    output_filename="SmartGrocery_AI_Frontend_Source_Code.docx",
    accent_color_hex=ACCENT_GREEN_HEX,
    accent_rgb=EMERALD_GREEN
)

print("\nAll split Word documents generated successfully!")
